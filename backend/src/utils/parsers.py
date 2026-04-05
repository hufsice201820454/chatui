"""
File content parsers (BE-FILE-02).
Supported: PDF, DOCX, XLSX, CSV, plain text.
Image (JPEG/PNG 등): 동기 파서는 None 반환. 실제 분석은 src.utils.vision_parser.parse_image_with_vision(비동기) 사용.
"""
import csv
import io
import logging
from pathlib import Path

logger = logging.getLogger("chatui.files.parsers")


def parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def parse_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def parse_xlsx(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"=== Sheet: {sheet.title} ===")
        for row in sheet.iter_rows(values_only=True):
            parts.append("\t".join(str(c) if c is not None else "" for c in row))
    return "\n".join(parts)


def parse_csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = ["\t".join(row) for row in reader]
    return "\n".join(rows)


def parse_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def parse_image(_data: bytes) -> str | None:
    """이미지는 동기 파싱 없음. Vision 분석은 vision_parser.parse_image_with_vision() 사용."""
    return None


PARSERS: dict[str, callable] = {
    "application/pdf": parse_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": parse_docx,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": parse_xlsx,
    "text/csv": parse_csv,
    "text/plain": parse_text,
    "application/octet-stream": parse_text,
    "image/jpeg": parse_image,
    "image/png": parse_image,
    "image/webp": parse_image,
}

# Extension fallback map
_EXT_MAP: dict[str, str] = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".txt": "text/plain",
    ".md": "text/plain",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
}


def parse_file(data: bytes, mime_type: str, filename: str = "") -> str | None:
    parser = PARSERS.get(mime_type)
    if not parser and filename:
        ext = Path(filename).suffix.lower()
        fallback_mime = _EXT_MAP.get(ext)
        if fallback_mime:
            parser = PARSERS.get(fallback_mime)

    if not parser:
        logger.warning("No parser for mime_type=%s filename=%s", mime_type, filename)
        return None

    try:
        return parser(data)
    except Exception as exc:
        logger.exception("Parsing failed for %s: %s", filename, exc)
        return None


SUPPORTED_MIME_TYPES: set[str] = set(PARSERS.keys())
SUPPORTED_EXTENSIONS: set[str] = set(_EXT_MAP.keys())

# Vision 파싱 대상 이미지 MIME (upload 시 parse_image_with_vision 호출)
IMAGE_MIME_TYPES: frozenset[str] = frozenset({"image/jpeg", "image/png", "image/webp"})
